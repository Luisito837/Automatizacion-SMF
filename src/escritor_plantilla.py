"""Copia las compras filtradas a la hoja de compras de la plantilla Excel.

Se modifica únicamente el XML de esa hoja dentro del archivo .xlsx. Todo lo demás
(tabla dinámica, segmentaciones, fórmulas, estilos) queda exactamente como estaba;
guardar con openpyxl, en cambio, borraría las segmentaciones de la plantilla.
"""
import datetime
import re
import xml.etree.ElementTree as ET
import zipfile
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import get_column_letter

from lector_excel import buscar_fila_encabezados
from utilidades import normalizar_texto

ENCABEZADOS_HOJA_COMPRAS = ("CódigoCompra", "FechaCompra", "CodProducto")

_ESPACIO_LIBRO = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_ESPACIO_RELACIONES = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_ATRIBUTO_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def parte_xml_de_la_hoja(ruta_libro, nombre_hoja) -> str:
    """Devuelve la ruta interna (dentro del .xlsx) del XML de una hoja, p. ej. 'xl/worksheets/sheet1.xml'."""
    with zipfile.ZipFile(ruta_libro) as archivo:
        libro = ET.fromstring(archivo.read("xl/workbook.xml"))
        relaciones = ET.fromstring(archivo.read("xl/_rels/workbook.xml.rels"))
    identificador = next(
        h.get(_ATRIBUTO_ID) for h in libro.iter(_ESPACIO_LIBRO + "sheet")
        if h.get("name") == nombre_hoja)
    destino = next(
        r.get("Target") for r in relaciones.iter(_ESPACIO_RELACIONES + "Relationship")
        if r.get("Id") == identificador)
    return destino.lstrip("/") if destino.startswith("/") else "xl/" + destino


class EscritorPlantilla:
    """Escribe en una COPIA de la plantilla; el archivo original no se modifica."""

    def volcar_compras(self, ruta_plantilla, filas, ruta_salida) -> int:
        """Devuelve cuántas filas se escribieron en la hoja de compras."""
        nombre_hoja, columnas = self._ubicar_hoja_de_compras(ruta_plantilla)
        parte_xml = self._parte_xml_de_la_hoja(ruta_plantilla, nombre_hoja)
        with zipfile.ZipFile(ruta_plantilla) as origen, \
                zipfile.ZipFile(ruta_salida, "w", zipfile.ZIP_DEFLATED) as destino:
            for parte in origen.infolist():
                contenido = origen.read(parte.filename)
                if parte.filename == parte_xml:
                    nuevo = self._reemplazar_datos(contenido.decode("utf-8"), columnas, filas)
                    contenido = nuevo.encode("utf-8")
                destino.writestr(parte, contenido)
        return len(filas)

    # ---------- ubicar la hoja ----------
    @staticmethod
    def _ubicar_hoja_de_compras(ruta_plantilla):
        """Devuelve (nombre de la hoja, {encabezado normalizado: número de columna})."""
        libro = openpyxl.load_workbook(ruta_plantilla, read_only=True)
        try:
            for hoja in libro.worksheets:
                hoja.reset_dimensions()
                primera = [tuple(f) for f in hoja.iter_rows(max_row=1, values_only=True)]
                ubicacion = buscar_fila_encabezados(primera, ENCABEZADOS_HOJA_COMPRAS)
                if ubicacion:
                    return hoja.title, {n: i + 1 for n, i in ubicacion[1].items()}
        finally:
            libro.close()
        raise ValueError(
            "La plantilla no tiene una hoja de compras (con las columnas "
            + ", ".join(ENCABEZADOS_HOJA_COMPRAS) + ").")

    @staticmethod
    def _parte_xml_de_la_hoja(ruta_plantilla, nombre_hoja) -> str:
        return parte_xml_de_la_hoja(ruta_plantilla, nombre_hoja)

    # ---------- reescribir los datos ----------
    def _reemplazar_datos(self, xml, columnas, filas) -> str:
        encabezado = re.search(r'<row r="1"(?:[^>]*/>|[^>]*>.*?</row>)', xml, re.S)
        if encabezado is None:
            raise ValueError("No se pudo leer la fila de encabezados de la hoja de compras.")
        filas_xml = "".join(
            self._xml_de_fila(numero, columnas, fila) for numero, fila in enumerate(filas, start=2))
        datos = f"<sheetData>{encabezado.group(0)}{filas_xml}</sheetData>"
        xml = re.sub(r"<sheetData>.*</sheetData>|<sheetData/>", lambda _: datos, xml, count=1, flags=re.S)
        dimension = f'<dimension ref="A1:{get_column_letter(max(columnas.values()))}{len(filas) + 1}"/>'
        return re.sub(r'<dimension ref="[^"]*"/>', lambda _: dimension, xml, count=1)

    def _xml_de_fila(self, numero, columnas, fila) -> str:
        """Escribe por nombre de columna: la columna RUC, que la plantilla no tiene, se omite sola."""
        celdas = []
        for encabezado, valor in fila.items():
            columna = columnas.get(normalizar_texto(encabezado))
            if columna:
                celdas.append((columna, self._xml_de_celda(f"{get_column_letter(columna)}{numero}", valor)))
        celdas.sort(key=lambda par: par[0])
        return f'<row r="{numero}">' + "".join(xml for _, xml in celdas) + "</row>"

    @staticmethod
    def _xml_de_celda(referencia, valor) -> str:
        if valor is None or valor == "":
            return ""
        if isinstance(valor, bool):
            return f'<c r="{referencia}" t="b"><v>{int(valor)}</v></c>'
        if isinstance(valor, (int, float)):
            return f'<c r="{referencia}"><v>{valor!r}</v></c>'
        if isinstance(valor, datetime.datetime):
            valor = valor.strftime("%d/%m/%Y %H:%M:%S")
        elif isinstance(valor, datetime.date):
            valor = valor.strftime("%d/%m/%Y")
        texto = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(valor))
        return f'<c r="{referencia}" t="inlineStr"><is><t xml:space="preserve">{escape(texto)}</t></is></c>'
