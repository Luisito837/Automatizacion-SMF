"""Exporta y recupera la distribución guardada dentro de la plantilla SMF."""
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl.utils import column_index_from_string, get_column_letter

from escritor_plantilla import parte_xml_de_la_hoja
from lector_excel import LectorExcel
from utilidades import normalizar_codigo, normalizar_texto

SEPARADOR = ", "

EQUIVALENCIAS_DE_CATEGORIA = {
    "COLAS NEGRAS": ("COLA NEGRA",),
    "COLAS BLANCAS": ("COLA BLANCA",),
    "SABORES": ("SAB",),
    "MALTA": ("MAL",),
    "JUGOS": ("JUG",),
    "TE": ("TE",),
    "BEBIDAS DEPORTIVAS": ("BEB.DEPOR.",),
    "TRADICIONAL": ("CER",),
    "LIGHT": ("CER",),
    "PREMIUN": ("CER",),
    "PREMIUM": ("CER",),
    "SANGRIA": ("SANG",),
    "VINO NACIONAL": ("VINOS",),
}

PATRON_ASIGNACION = re.compile(
    r"(?P<codigo>[A-Za-z0-9_-]+)\s+(?P<cantidad>\d+)\s*(?P<modo>UND|UN|CJ|CAJA|CAJAS|BTO|BULTO|BULTOS)\b",
    re.IGNORECASE,
)


@dataclass
class ResultadoExportacion:
    ruta_salida: Path
    clientes_escritos: int = 0
    entradas_escritas: int = 0
    productos_sin_columna: list = field(default_factory=list)
    clientes_sin_fila: list = field(default_factory=list)

    def resumen(self) -> str:
        lineas = [
            f"Clientes con datos exportados: {self.clientes_escritos}",
            f"Asignaciones exportadas: {self.entradas_escritas}",
            "",
            f"Archivo:\n{self.ruta_salida}",
        ]
        if self.productos_sin_columna:
            lineas += ["", "No se exportaron (la plantilla no tiene columna para su categoría):",
                       *[f"   · {nombre}" for nombre in self.productos_sin_columna]]
        if self.clientes_sin_fila:
            lineas += ["", "Clientes que no se encontraron en la hoja de clientes:",
                       *[f"   · {nombre}" for nombre in self.clientes_sin_fila]]
        return "\n".join(lineas)


class ExportadorPlantilla:
    def __init__(self, lector_excel: LectorExcel):
        self.lector_excel = lector_excel

    def exportar(self, ruta_base, asignaciones, ruta_salida) -> ResultadoExportacion:
        """Copia el libro, limpia la distribución previa y escribe el estado actual completo."""
        ruta_base, ruta_salida = Path(ruta_base), Path(ruta_salida)
        if ruta_base.resolve() == ruta_salida.resolve():
            raise ValueError("Elija un nombre distinto al del archivo de origen para no sobrescribirlo.")
        hoja = self.lector_excel.buscar_hoja_con_encabezados(ruta_base, ("DIA", "CLIENTE"))
        if hoja is None:
            raise ValueError("La plantilla no tiene una hoja de clientes (con las columnas DIA y CLIENTE).")

        filas_de_clientes = self._filas_de_clientes(hoja)
        columnas = self._columnas_de_categoria(hoja)
        resultado = ResultadoExportacion(ruta_salida)
        celdas = self._agrupar_en_celdas(asignaciones, filas_de_clientes, columnas, resultado)
        filas_limpiables = set(filas_de_clientes.values())
        columnas_limpiables = {columna for columna, _es_producto in columnas.values()}

        parte_xml = parte_xml_de_la_hoja(ruta_base, hoja.nombre)
        with zipfile.ZipFile(ruta_base) as origen, zipfile.ZipFile(ruta_salida, "w", zipfile.ZIP_DEFLATED) as destino:
            for parte in origen.infolist():
                contenido = origen.read(parte.filename)
                if parte.filename == parte_xml:
                    contenido = self._escribir_celdas(
                        contenido.decode("utf-8"), celdas, filas_limpiables, columnas_limpiables
                    ).encode("utf-8")
                destino.writestr(parte, contenido)
        return resultado

    def leer_asignaciones_guardadas(self, ruta, productos) -> list:
        """Lee cadenas de distribución exportadas, para reabrir una plantilla como un guardado."""
        hoja = self.lector_excel.buscar_hoja_con_encabezados(ruta, ("DIA", "CLIENTE"))
        if hoja is None:
            return []
        col = hoja.columnas
        idx_dia, idx_cliente = col["DIA"], col["CLIENTE"]
        idx_codigo = col.get("CODIGO SISTEMA")
        idx_cep = col.get("CEP")
        producto_por_codigo = {normalizar_codigo(p.codigo): p for p in productos}
        registros = []
        for indice_fila, fila in enumerate(hoja.filas):
            if indice_fila <= hoja.fila_encabezados:
                continue
            nombre = self._valor(fila, idx_cliente)
            dia = self._valor(fila, idx_dia)
            if not nombre or not dia:
                continue
            codigo_cliente = self._valor(fila, idx_codigo) if idx_codigo is not None else ""
            cep = self._valor(fila, idx_cep) if idx_cep is not None else ""
            for valor in fila:
                if not isinstance(valor, str) or not valor.strip():
                    continue
                for coincidencia in PATRON_ASIGNACION.finditer(valor):
                    codigo = normalizar_codigo(coincidencia.group("codigo"))
                    if codigo not in producto_por_codigo:
                        continue
                    modo_texto = normalizar_texto(coincidencia.group("modo"))
                    modo = "BULTO" if modo_texto.startswith(("BTO", "BULTO")) else (
                        "UNIDAD" if modo_texto.startswith(("UND", "UN")) else "CAJA")
                    registros.append({
                        "dia": dia,
                        "cliente": nombre,
                        "codigo_sistema": codigo_cliente,
                        "cep": cep,
                        "producto_codigo": codigo,
                        "cantidad": int(coincidencia.group("cantidad")),
                        "modo": modo,
                    })
        return registros

    # ---------- localizar clientes y columnas ----------
    @staticmethod
    def _filas_de_clientes(hoja) -> dict:
        """Índice de fila por día+código/CEP/nombre para no confundir clientes homónimos."""
        col = hoja.columnas
        idx_dia, idx_cliente = col["DIA"], col["CLIENTE"]
        idx_codigo, idx_cep = col.get("CODIGO SISTEMA"), col.get("CEP")
        filas = {}
        for numero, fila in enumerate(hoja.filas):
            if numero <= hoja.fila_encabezados:
                continue
            dia = normalizar_texto(ExportadorPlantilla._valor(fila, idx_dia))
            nombre = normalizar_texto(ExportadorPlantilla._valor(fila, idx_cliente))
            if not dia or not nombre:
                continue
            excel_row = numero + 1
            codigo = normalizar_texto(ExportadorPlantilla._valor(fila, idx_codigo)) if idx_codigo is not None else ""
            cep = normalizar_texto(ExportadorPlantilla._valor(fila, idx_cep)) if idx_cep is not None else ""
            if codigo:
                filas[(dia, "CODIGO", codigo)] = excel_row
            if cep:
                filas[(dia, "CEP", cep)] = excel_row
            filas[(dia, "NOMBRE", nombre)] = excel_row
        return filas

    @staticmethod
    def _columnas_de_categoria(hoja) -> dict:
        """Lee etiquetas de categoría de la fila justo encima de DIA/CLIENTE."""
        if hoja.fila_encabezados == 0:
            return {}
        columnas = {}
        for indice, valor in enumerate(hoja.filas[hoja.fila_encabezados - 1]):
            texto = normalizar_texto(valor)
            if texto:
                columnas[texto.lstrip("%").strip()] = (indice + 1, not texto.startswith("%"))
        return columnas

    @staticmethod
    def _columna_para(producto, columnas):
        descripcion = normalizar_texto(producto.descripcion)
        for texto, (numero, es_producto) in columnas.items():
            if es_producto and texto in descripcion:
                return numero
        categoria = normalizar_texto(producto.categoria)
        for candidato in (*EQUIVALENCIAS_DE_CATEGORIA.get(categoria, ()), categoria):
            if candidato in columnas and not columnas[candidato][1]:
                return columnas[candidato][0]
        return None

    def _agrupar_en_celdas(self, asignaciones, filas_de_clientes, columnas, resultado) -> dict:
        textos = {}
        clientes_escritos, sin_fila, sin_columna = set(), set(), set()
        for asignacion in asignaciones:
            cliente = asignacion.cliente
            dia = normalizar_texto(cliente.dia)
            codigo = normalizar_texto(getattr(cliente, "codigo_sistema", ""))
            cep = normalizar_texto(getattr(cliente, "cep", ""))
            nombre = normalizar_texto(cliente.nombre)
            fila = (filas_de_clientes.get((dia, "CODIGO", codigo)) if codigo else None)
            fila = fila or (filas_de_clientes.get((dia, "CEP", cep)) if cep else None)
            fila = fila or filas_de_clientes.get((dia, "NOMBRE", nombre))
            if fila is None:
                sin_fila.add(cliente.nombre)
                continue
            columna = self._columna_para(asignacion.producto, columnas)
            if columna is None:
                sin_columna.add(asignacion.producto.descripcion)
                continue
            textos.setdefault((fila, columna), []).append(asignacion.texto())
            clientes_escritos.add(fila)
            resultado.entradas_escritas += 1
        resultado.clientes_escritos = len(clientes_escritos)
        resultado.productos_sin_columna = sorted(sin_columna)
        resultado.clientes_sin_fila = sorted(sin_fila)
        celdas = {}
        for (fila, columna), lista in textos.items():
            celdas.setdefault(fila, {})[columna] = SEPARADOR.join(lista)
        return celdas

    # ---------- XML interno de la hoja de clientes ----------
    @staticmethod
    def _escribir_celdas(xml, celdas_por_fila, filas_a_limpiar=None, columnas_a_limpiar=None) -> str:
        filas_a_limpiar = filas_a_limpiar or set()
        columnas_a_limpiar = columnas_a_limpiar or set()
        patron_fila = re.compile(r'<row r="(\d+)"([^>]*?)(?:/>|>(.*?)</row>)', re.S)
        patron_celda = re.compile(r'<c r="([A-Z]+)\d+"[^>]*?(?:/>|>.*?</c>)', re.S)

        def reemplazar(coincidencia):
            fila = int(coincidencia.group(1))
            nuevas = celdas_por_fila.get(fila, {})
            if fila not in filas_a_limpiar and not nuevas:
                return coincidencia.group(0)
            atributos = re.sub(r'\sspans="[^"]*"', "", coincidencia.group(2))
            celdas = {column_index_from_string(c.group(1)): c.group(0)
                      for c in patron_celda.finditer(coincidencia.group(3) or "")}
            if fila in filas_a_limpiar:
                for columna in columnas_a_limpiar:
                    if columna in celdas:
                        anterior = celdas[columna]
                        estilo = re.search(r'\ss="(\d+)"', anterior.split(">")[0])
                        atributo_estilo = f' s="{estilo.group(1)}"' if estilo else ""
                        celdas[columna] = f'<c r="{get_column_letter(columna)}{fila}"{atributo_estilo}/>'
            for columna, texto in nuevas.items():
                anterior = celdas.get(columna, "")
                estilo = re.search(r'\ss="(\d+)"', anterior.split(">")[0])
                atributo_estilo = f' s="{estilo.group(1)}"' if estilo else ""
                limpio = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", texto)
                celdas[columna] = (
                    f'<c r="{get_column_letter(columna)}{fila}"{atributo_estilo} t="inlineStr">'
                    f'<is><t xml:space="preserve">{escape(limpio)}</t></is></c>')
            return f'<row r="{fila}"{atributos}>' + "".join(celdas[k] for k in sorted(celdas)) + "</row>"

        xml = patron_fila.sub(reemplazar, xml)
        return xml

    @staticmethod
    def _valor(fila, indice) -> str:
        if indice is None or indice < 0 or indice >= len(fila) or fila[indice] is None:
            return ""
        return str(fila[indice]).strip()
