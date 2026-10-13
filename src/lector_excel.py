"""Lectura de archivos Excel: reporte general de compras y hojas de apoyo."""
import warnings
from dataclasses import dataclass

import openpyxl

from utilidades import normalizar_texto

warnings.filterwarnings("ignore", message="Unknown extension")

ENCABEZADOS_REPORTE = (
    "CodProducto", "DescrProducto", "Cantidad", "UnidadCompra",
    "DescrCategoría", "CodTipoInv", "Comentario Cabecera",
)


def buscar_fila_encabezados(filas, obligatorios, maximo_filas=40):
    """Busca entre las primeras filas la que contiene todos los encabezados pedidos.

    Devuelve (numero_de_fila, {encabezado_normalizado: indice_de_columna}) o None.
    """
    requeridos = {normalizar_texto(nombre) for nombre in obligatorios}
    for numero, fila in enumerate(filas[:maximo_filas]):
        columnas = {}
        for indice, valor in enumerate(fila):
            nombre = normalizar_texto(valor)
            if nombre:
                columnas[nombre] = indice
        if requeridos <= set(columnas):
            return numero, columnas
    return None


@dataclass
class HojaLeida:
    nombre: str
    filas: list
    fila_encabezados: int
    columnas: dict


@dataclass
class TablaCompras:
    encabezados: list
    filas: list  # cada fila es un diccionario {encabezado: valor}


class LectorExcel:
    """Abre un libro de Excel y localiza la hoja que tenga ciertos encabezados."""

    def buscar_hoja_con_encabezados(self, ruta, obligatorios):
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        try:
            for hoja in libro.worksheets:
                # Algunos reportes declaran mal sus dimensiones y se perderían columnas.
                hoja.reset_dimensions()
                muestra = [tuple(f) for f in hoja.iter_rows(max_row=40, values_only=True)]
                ubicacion = buscar_fila_encabezados(muestra, obligatorios)
                if ubicacion:
                    todas = [tuple(f) for f in hoja.iter_rows(values_only=True)]
                    return HojaLeida(hoja.title, todas, ubicacion[0], ubicacion[1])
        finally:
            libro.close()
        return None


class LectorReporteCompras:
    """Convierte el Reporte General de Compras en una tabla de filas con nombre."""

    def __init__(self, lector_excel: LectorExcel):
        self.lector_excel = lector_excel

    def leer(self, ruta) -> TablaCompras:
        hoja = self.lector_excel.buscar_hoja_con_encabezados(ruta, ENCABEZADOS_REPORTE)
        if hoja is None:
            raise ValueError(
                "El archivo no parece un Reporte General de Compras.\n"
                "No se encontraron las columnas: " + ", ".join(ENCABEZADOS_REPORTE)
            )
        encabezados = hoja.filas[hoja.fila_encabezados]
        posiciones = [
            (indice, str(valor).strip())
            for indice, valor in enumerate(encabezados)
            if valor is not None and str(valor).strip()
        ]
        filas = []
        for fila in hoja.filas[hoja.fila_encabezados + 1:]:
            if all(valor is None or str(valor).strip() == "" for valor in fila):
                continue
            filas.append({
                nombre: (fila[indice] if indice < len(fila) else None)
                for indice, nombre in posiciones
            })
        return TablaCompras([nombre for _, nombre in posiciones], filas)
